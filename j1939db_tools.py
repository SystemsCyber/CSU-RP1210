"""
Tools for building and checking CSU-RP1210 J1939 databases.

* generate():  SAE J1939 Digital Annex (.xlsx/.xls) -> J1939db.licensed.json (metric)
               and J1939db.us.licensed.json (US customary), using pretty_j1939's
               Digital Annex converter, then adding the legacy fields the Python
               application reads (StartBit, EndBit, ...).
* validate():  structural and consistency checks for any J1939db*.json version.
* compare():   differences between two database versions; cross-checks the unit
               conversion when one is metric and the other US customary.
* run_vectors(): decode test vectors (PGN + data -> expected SPN values in each
               unit system) against a database.

The same functions back the DigitalAnnexSelect dialog, the pytest suite and the
command line:

    python j1939db_tools.py generate DA.xlsx [--out DIR]
    python j1939db_tools.py validate J1939db.licensed.json [--baseline OLD.json]
                                     [--vectors tests/j1939db_vectors.json]
"""

import argparse
import contextlib
import copy
import datetime
import hashlib
import io
import json
import os
import re
import sys
import tempfile

MODULE_DIR = os.path.dirname(os.path.abspath(__file__))
UNITS_FILE = os.path.join(MODULE_DIR, "j1939_units.json")
VECTORS_FILE = os.path.join(MODULE_DIR, "tests", "j1939db_vectors.json")
SETTINGS_FILE = "csu_settings.json"

METRIC = "metric"
US = "us"
OUTPUT_NAMES = {METRIC: "J1939db.licensed.json", US: "J1939db.us.licensed.json"}

PASS, WARN, FAIL, INFO = "PASS", "WARN", "FAIL", "INFO"
SPN_STATUSES = ("Valid", "NotAvailable", "Error", "Reserved", "Missing")


# --------------------------------------------------------------------------
# Settings (shared with the Rust core: csu_settings.json, key "j1939_units")
# --------------------------------------------------------------------------

def read_unit_preference(directory=None):
    """Preferred unit system: $CSU_UNITS, then csu_settings.json, else metric."""
    env = os.environ.get("CSU_UNITS", "").strip().lower()
    if env in (METRIC, US):
        return env
    path = os.path.join(directory or os.getcwd(), SETTINGS_FILE)
    try:
        with open(path) as f:
            value = str(json.load(f).get("j1939_units", METRIC)).lower()
        return value if value in (METRIC, US) else METRIC
    except (OSError, ValueError, AttributeError):
        return METRIC


def write_unit_preference(units, directory=None):
    path = os.path.join(directory or os.getcwd(), SETTINGS_FILE)
    settings = {}
    try:
        with open(path) as f:
            settings = json.load(f)
    except (OSError, ValueError):
        pass
    settings["j1939_units"] = units
    with open(path, "w") as f:
        json.dump(settings, f, indent=2)
    return path


def database_candidates(directory, units):
    """Licensed database file names in search order for a unit preference."""
    first, second = (OUTPUT_NAMES[US], OUTPUT_NAMES[METRIC]) if units == US else (OUTPUT_NAMES[METRIC], OUTPUT_NAMES[US])
    return [os.path.join(directory, n) for n in (first, second, "J1939db.json")]


# --------------------------------------------------------------------------
# Units
# --------------------------------------------------------------------------

def canonical_unit(unit):
    """Normalize a unit string for table lookup ('°C' -> 'degc', 'N·m' -> 'nm')."""
    s = str(unit or "").strip().lower()
    s = s.replace("°", "deg").replace("degrees", "deg").replace("µ", "u").replace("μ", "u")
    s = s.translate(str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹⁻", "0123456789-"))
    s = re.sub(r"[\s.\-_*·⋅•^]", "", s)
    return s


class UnitTable:
    def __init__(self, path=UNITS_FILE):
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        self.conversions = {}
        self.by_alias = {}
        for key, conv in data["conversions"].items():
            self.conversions[key] = conv
            for alias in conv.get("aliases", []) + [key]:
                self.by_alias[canonical_unit(alias)] = key
        self.labels = {k: v for k, v in data.get("labels", {}).items() if not k.startswith("_")}
        self.us_labels = {canonical_unit(c["us"]): k for k, c in self.conversions.items()}

    def lookup(self, unit):
        """Conversion entry for a metric unit, or None."""
        key = self.by_alias.get(canonical_unit(unit))
        return self.conversions.get(key) if key else None

    def metric_label(self, unit):
        conv = self.lookup(unit)
        if conv:
            return conv["metric"]
        return self.labels.get(canonical_unit(unit), unit)

    def system_of(self, unit):
        """'metric', 'us' or None (unit-less or not convertible)."""
        c = canonical_unit(unit)
        if c in self.by_alias:
            return METRIC
        if c in self.us_labels:
            return US
        return None

    def conversion_between(self, metric_unit, us_unit):
        conv = self.lookup(metric_unit)
        if conv and canonical_unit(conv["us"]) == canonical_unit(us_unit):
            return conv
        return None


def _num(x):
    if isinstance(x, bool):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return float(str(x).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def _clean(x):
    """Round away binary float noise (12 significant digits)."""
    return float(f"{x:.12g}")


_RANGE_RE = re.compile(r"^\s*(-?[\d,]*\.?\d+(?:[eE][-+]?\d+)?)\s*to\s*(-?[\d,]*\.?\d+(?:[eE][-+]?\d+)?)\s*(.*?)\s*$")


def _convert_range_text(text, conv, label):
    m = _RANGE_RE.match(str(text or ""))
    if not m:
        return text
    lo, hi = (_num(m.group(1)), _num(m.group(2)))
    if lo is None or hi is None:
        return text
    lo, hi = lo * conv["scale"] + conv["offset"], hi * conv["scale"] + conv["offset"]
    return f"{lo:.6g} to {hi:.6g} {label}"


# --------------------------------------------------------------------------
# Start bits (pretty_j1939 schema)
# --------------------------------------------------------------------------

def resolve_start_bit(entry, length):
    """Resolve a SPNStartBits entry to one contiguous start bit, or None.

    pretty_j1939 stores byte ranges as [first, last-byte-start] (e.g. "1-4" ->
    [0, 24]). A pair is one contiguous little-endian field when its last bit
    falls in the byte that starts at the second value.
    """
    if isinstance(entry, (int, float)):
        return int(entry) if entry >= 0 else None
    if not isinstance(entry, list) or not entry:
        return None
    vals = [int(v) for v in entry if _num(v) is not None]
    if not vals or vals[0] < 0:
        return None
    if len(vals) == 1:
        return vals[0]
    if len(vals) == 2 and isinstance(length, int) and vals[1] >= vals[0]:
        end = vals[0] + length - 1
        if vals[1] <= end < vals[1] + 8:
            return vals[0]
    return None


def spn_start_in_pgn(db, pgn_key, index):
    pgn = db["J1939PGNdb"][pgn_key]
    spn_key = str(pgn["SPNs"][index])
    spn = db["J1939SPNdb"].get(spn_key, {})
    length = spn.get("SPNLength")
    starts = pgn.get("SPNStartBits")
    if starts is not None and index < len(starts):
        return resolve_start_bit(starts[index], length if isinstance(length, int) else None)
    sb = _num(spn.get("StartBit"))
    return int(sb) if sb is not None and sb >= 0 else None


# --------------------------------------------------------------------------
# Generation
# --------------------------------------------------------------------------

def _sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def convert_digital_annex(da_paths, log=None):
    """Run pretty_j1939's converter and return the raw database dict."""
    try:
        from pretty_j1939.create_j1939db_json import J1939daConverter
    except ImportError as e:
        raise RuntimeError("pretty_j1939 is required: python -m pip install pretty_j1939") from e
    with tempfile.TemporaryDirectory() as tmp:
        out = os.path.join(tmp, "j1939db.json")
        captured = io.StringIO()
        with contextlib.redirect_stderr(captured), contextlib.redirect_stdout(captured):
            J1939daConverter(list(da_paths)).convert(out)
        if log:
            for line in captured.getvalue().splitlines():
                if line.strip():
                    log(line)
        with open(out, encoding="utf-8") as f:
            return json.load(f)


def add_legacy_fields(db, table):
    """Make a pretty_j1939 database readable by the CSU-RP1210 Python app.

    Adds per-SPN StartBit/EndBit/SPN/Acronym (from the first PGN that carries
    the SPN), upper-cases ASCII units, gives raw 'binary' fields resolution 1,
    applies display labels and ensures the supplementary tables exist.
    """
    spns = db.setdefault("J1939SPNdb", {})
    pgns = db.setdefault("J1939PGNdb", {})
    for pgn_key, pgn in pgns.items():
        for i, spn_num in enumerate(pgn.get("SPNs", [])):
            spn = spns.get(str(spn_num))
            if spn is None or "StartBit" in spn:
                continue
            start = spn_start_in_pgn(db, pgn_key, i)
            length = spn.get("SPNLength")
            spn["StartBit"] = start if start is not None else -1
            spn["EndBit"] = start + length - 1 if (start is not None and isinstance(length, int)) else -1
            spn["Acronym"] = pgn.get("Label", "")
            spn["PGNLength"] = pgn.get("PGNLength", "")
            spn["TransmissionRate"] = pgn.get("Rate", "")
    for num, spn in spns.items():
        spn.setdefault("SPN", int(num))
        spn.setdefault("StartBit", -1)
        units = spn.get("Units", "")
        if canonical_unit(units) == "ascii":
            spn["Units"] = "ASCII"
        elif canonical_unit(units) == "binary" and not _num(spn.get("Resolution")):
            spn["Resolution"] = 1
        else:
            spn["Units"] = table.metric_label(units)
    for name in ("J1939BitDecodings", "J1939SATabledb", "J1939SAHWTabledb", "J1939FMITabledb",
                 "J1939LampFlashTabledb", "J1939OBDTabledb", "J1939Manufacturerdb"):
        db.setdefault(name, {})
    return db


def apply_value_only_scaling(db, da_paths, log=None):
    """Prefer the Digital Annex numeric 'Scale Factor / Offset (value only)' columns.

    pretty_j1939 derives Resolution and Offset by parsing the text columns, which
    mis-reads units that contain digits (e.g. "0.1 m/s2 per bit" -> 0.05). For SPNs
    whose SLOT has a numeric transfer function, the value-only columns are
    authoritative. Returns the list of corrected SPNs.
    """
    try:
        slots, spn_slot, spn_row = read_slots(da_paths)
    except Exception as e:  # older workbooks without value-only columns
        if log:
            log(f"Value-only scaling not applied: {e}")
        return []
    corrected = []
    for num, spn in db.get("J1939SPNdb", {}).items():
        slot, row = slots.get(spn_slot.get(num)), spn_row.get(num)
        if not slot or slot["transfer"] != "numeric" or not row or row["scale"] is None:
            continue
        res, off = _num(spn.get("Resolution")), _num(spn.get("Offset")) or 0.0
        new_off = row["offset"] if row["offset"] is not None else off
        if not (_close(res, row["scale"], 1e-9) and _close(off, new_off, 1e-9)):
            corrected.append({"spn": int(num), "parsed": [res, off], "value_only": [row["scale"], new_off]})
            spn["Resolution"], spn["Offset"] = row["scale"], new_off
    if log and corrected:
        log(f"Corrected scaling of {len(corrected)} SPNs from the Digital Annex value-only columns")
    return corrected


def to_us_customary(metric_db, table):
    """Return a US customary copy of a metric database plus conversion counts."""
    db = copy.deepcopy(metric_db)
    counts = {}
    for spn in db["J1939SPNdb"].values():
        conv = table.lookup(spn.get("Units", ""))
        if not conv:
            continue
        scale, off = conv["scale"], conv["offset"]
        res, offset = _num(spn.get("Resolution")), _num(spn.get("Offset")) or 0.0
        if res:
            spn["Resolution"] = _clean(res * scale)
        spn["Offset"] = _clean(offset * scale + off)
        for k in ("OperationalLow", "OperationalHigh"):
            v = _num(spn.get(k))
            if v is not None:
                spn[k] = _clean(v * scale + off)
        spn["DataRange"] = _convert_range_text(spn.get("DataRange"), conv, conv["us"])
        spn["OperationalRange"] = _convert_range_text(spn.get("OperationalRange"), conv, conv["us"])
        spn["MetricUnits"] = conv["metric"]
        spn["Units"] = conv["us"]
        key = f'{conv["metric"]} -> {conv["us"]}'
        counts[key] = counts.get(key, 0) + 1
    return db, counts


def generate(da_paths, out_dir, systems=(METRIC, US), units_path=UNITS_FILE, log=print):
    """Build licensed databases from Digital Annex files. Returns {system: path}."""
    table = UnitTable(units_path)
    log(f"Reading {len(da_paths)} Digital Annex file(s)...")
    raw = convert_digital_annex(da_paths, log=log)
    if not raw.get("J1939PGNdb") or not raw.get("J1939SPNdb"):
        raise ValueError("No PGN/SPN data found. Is this a J1939 Digital Annex workbook "
                         "(sheet 'SPs & PGs' or 'SPNs & PGNs')?")
    metric = add_legacy_fields(raw, table)
    corrections = apply_value_only_scaling(metric, da_paths, log)
    try:
        from importlib.metadata import version
        generator = f"pretty_j1939 {version('pretty_j1939')}"
    except Exception:
        generator = "pretty_j1939"
    base_meta = {
        "skeleton": False,
        "generated_by": "CSU-RP1210 DigitalAnnexSelect",
        "generator": generator,
        "generated_at": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
        "sources": [{"file": os.path.basename(p), "sha256": _sha256(p)} for p in da_paths],
        "schema": "current (PGN SPNStartBits) with legacy per-SPN StartBit fields",
        "scaling_corrections": corrections,
        "license_notice": "Derived from the SAE J1939 Digital Annex under the user's license. Do not redistribute or commit.",
    }
    outputs = {}
    us_db, counts = to_us_customary(metric, table)
    variants = {METRIC: (metric, {}), US: (us_db, counts)}
    os.makedirs(out_dir, exist_ok=True)
    for system in systems:
        db, conv_counts = variants[system]
        db = dict(db)
        db["_meta"] = dict(base_meta, units=system, unit_conversions=conv_counts)
        path = os.path.join(out_dir, OUTPUT_NAMES[system])
        with open(path, "w", encoding="utf-8") as f:
            json.dump(db, f, indent=1)
        outputs[system] = path
        log(f"Wrote {path}: {len(db['J1939PGNdb'])} PGNs, {len(db['J1939SPNdb'])} SPNs"
            + (f", {sum(conv_counts.values())} SPNs converted to US units" if system == US else ""))
    return outputs


# --------------------------------------------------------------------------
# Decoding (reference implementation used by the vectors)
# --------------------------------------------------------------------------

def extract_bits(data, start, length):
    if length <= 0 or length > 64 or start < 0 or start + length > len(data) * 8:
        return None
    value = int.from_bytes(bytes(data), "little")
    return (value >> start) & ((1 << length) - 1)


def classify(raw, length):
    if length < 2:
        return "Valid"
    if length < 8:
        full = (1 << length) - 1
        return "NotAvailable" if raw == full else "Error" if raw == full - 1 else "Valid"
    top = (raw >> (length - 8)) & 0xFF
    return {0xFF: "NotAvailable", 0xFE: "Error"}.get(top, "Reserved" if top >= 0xFB else "Valid")


def decode_spn(db, pgn, spn, data):
    """Decode one SPN. Returns dict(value, raw, units, status) or None if undefined."""
    pgn_entry = db.get("J1939PGNdb", {}).get(str(pgn))
    spn_entry = db.get("J1939SPNdb", {}).get(str(spn))
    if pgn_entry is None or spn_entry is None or spn not in pgn_entry.get("SPNs", []):
        return None
    index = pgn_entry["SPNs"].index(spn)
    start = spn_start_in_pgn(db, str(pgn), index)
    length = spn_entry.get("SPNLength")
    units = spn_entry.get("Units", "")
    result = {"value": None, "raw": None, "units": units, "status": "Missing", "text": None}
    if canonical_unit(units) == "ascii":
        # '*'-delimited text fields, assigned in SPN order (e.g. VIN, component ID).
        ascii_spns = [s for s in pgn_entry["SPNs"]
                      if canonical_unit(db["J1939SPNdb"].get(str(s), {}).get("Units")) == "ascii"]
        first = ascii_spns[0]
        offset = spn_start_in_pgn(db, str(pgn), pgn_entry["SPNs"].index(first))
        fields = bytes(data)[(offset or 0) // 8:].split(b"*")
        position = ascii_spns.index(spn)
        if position < len(fields) and fields[position]:
            result["text"] = fields[position].decode("ascii", "replace").strip()
            result["status"] = "Valid"
        return result
    if start is None or not isinstance(length, int):
        return result
    raw = extract_bits(data, start, length)
    if raw is None:
        return result
    result["raw"] = raw
    result["status"] = classify(raw, length)
    res = _num(spn_entry.get("Resolution")) or 0.0
    if res <= 0:
        res = 1.0
    result["value"] = _clean(raw * res + (_num(spn_entry.get("Offset")) or 0.0))
    return result


# --------------------------------------------------------------------------
# Validation
# --------------------------------------------------------------------------

class Check:
    def __init__(self, name, level, message, details=None):
        self.name, self.level, self.message = name, level, message
        self.details = list(details or [])

    def __repr__(self):
        return f"{self.level:4} {self.name}: {self.message}"


def load(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def unit_system(db, table=None):
    """Declared (_meta.units) or inferred unit system of a database."""
    declared = db.get("_meta", {}).get("units")
    if declared in (METRIC, US):
        return declared
    table = table or UnitTable()
    votes = {METRIC: 0, US: 0}
    for spn in db.get("J1939SPNdb", {}).values():
        s = table.system_of(spn.get("Units"))
        if s:
            votes[s] += 1
    if not any(votes.values()):
        return None
    return US if votes[US] > votes[METRIC] else METRIC


def _sample(items, n=12):
    items = list(items)
    return items[:n] + ([f"... {len(items) - n} more"] if len(items) > n else [])


def validate(db, table=None):
    table = table or UnitTable()
    checks = []
    meta = db.get("_meta", {})
    pgns, spns = db.get("J1939PGNdb"), db.get("J1939SPNdb")

    if not isinstance(pgns, dict) or not isinstance(spns, dict):
        return [Check("Structure", FAIL, "J1939PGNdb and J1939SPNdb tables are required")]
    checks.append(Check("Structure", PASS, f"{len(pgns)} PGNs, {len(spns)} SPNs, "
                        f"{len(db.get('J1939BitDecodings', {}))} bit decodings, "
                        f"{len(db.get('J1939SATabledb', {}))} source addresses"))

    if meta.get("skeleton"):
        checks.append(Check("Content", WARN, "Skeleton database: no licensed Digital Annex content"))
    elif len(pgns) < 500 or len(spns) < 2000:
        checks.append(Check("Content", WARN, "Fewer PGNs/SPNs than a complete Digital Annex "
                            "(expected at least 500 PGNs and 2000 SPNs)"))
    else:
        checks.append(Check("Content", PASS, "Size is consistent with a complete Digital Annex"))

    if meta.get("sources"):
        src = ", ".join(f'{s.get("file")} (sha256 {str(s.get("sha256"))[:12]}...)' for s in meta["sources"])
        checks.append(Check("Provenance", INFO, f'{meta.get("generator", "")} on {meta.get("generated_at", "?")} from {src}'))

    missing, length_mismatch, undecodable, overflow, overlaps = [], [], [], [], []
    for key, pgn in pgns.items():
        spn_list = pgn.get("SPNs", [])
        starts = pgn.get("SPNStartBits")
        if starts is not None and len(starts) != len(spn_list):
            length_mismatch.append(f"PGN {key}: {len(spn_list)} SPNs vs {len(starts)} start bits")
        pgn_len = _num(pgn.get("PGNLength"))
        limit = int(pgn_len) * 8 if pgn_len else None
        fields = []
        for i, s in enumerate(spn_list):
            spn = spns.get(str(s))
            if spn is None:
                missing.append(f"PGN {key} -> SPN {s}")
                continue
            length = spn.get("SPNLength")
            if not isinstance(length, int):
                continue
            start = spn_start_in_pgn(db, key, i)
            if start is None:
                if canonical_unit(spn.get("Units")) != "ascii":
                    undecodable.append(f"PGN {key} SPN {s}")
                continue
            if limit and start + length > limit:
                overflow.append(f"PGN {key} SPN {s}: bits {start}..{start + length - 1} beyond {limit} bits")
            fields.append((start, start + length, s))
        fields.sort()
        for (a0, a1, sa), (b0, b1, sb) in zip(fields, fields[1:]):
            if b0 < a1:
                overlaps.append(f"PGN {key}: SPN {sa} and SPN {sb}")
    checks.append(Check("SPN references", FAIL if missing else PASS,
                        f"{len(missing)} PGN entries reference undefined SPNs" if missing else "Every PGN SPN is defined",
                        _sample(missing)))
    checks.append(Check("Start-bit lists", FAIL if length_mismatch else PASS,
                        f"{len(length_mismatch)} PGNs with mismatched SPNStartBits" if length_mismatch else "SPNStartBits align with SPNs",
                        _sample(length_mismatch)))
    checks.append(Check("Decodable positions", WARN if undecodable else PASS,
                        f"{len(undecodable)} non-ASCII SPNs have unknown or split positions" if undecodable else "All fixed-position SPNs resolve",
                        _sample(undecodable)))
    checks.append(Check("Field bounds", WARN if overflow else PASS,
                        f"{len(overflow)} SPNs extend past their PGN length" if overflow else "All fields fit their PGN length",
                        _sample(overflow)))
    checks.append(Check("Overlapping fields", WARN if overlaps else PASS,
                        f"{len(overlaps)} overlapping SPN pairs (may be intentional multiplexing)" if overlaps else "No overlapping SPNs",
                        _sample(overlaps)))

    bad_numbers, bad_ranges, no_legacy = [], [], []
    for num, spn in spns.items():
        length = spn.get("SPNLength")
        if not (isinstance(length, int) and length > 0) and not str(length).startswith("Variable"):
            bad_numbers.append(f"SPN {num}: SPNLength {length!r}")
        if spn.get("Resolution") is not None and _num(spn.get("Resolution")) is None:
            bad_numbers.append(f"SPN {num}: Resolution {spn.get('Resolution')!r}")
        lo, hi = _num(spn.get("OperationalLow")), _num(spn.get("OperationalHigh"))
        if lo is not None and hi is not None and lo > hi:
            bad_ranges.append(f"SPN {num}: {lo} > {hi}")
        if "StartBit" not in spn:
            no_legacy.append(f"SPN {num}")
    checks.append(Check("Numeric fields", FAIL if bad_numbers else PASS,
                        f"{len(bad_numbers)} malformed numeric fields" if bad_numbers else "Lengths and resolutions are numeric",
                        _sample(bad_numbers)))
    checks.append(Check("Operational ranges", WARN if bad_ranges else PASS,
                        f"{len(bad_ranges)} SPNs with low > high" if bad_ranges else "Operational ranges are ordered",
                        _sample(bad_ranges)))
    checks.append(Check("Python app compatibility", WARN if no_legacy else PASS,
                        f"{len(no_legacy)} SPNs lack StartBit (the CSU-RP1210 Python tabs will fail on them; "
                        "regenerate with DigitalAnnexSelect)" if no_legacy else "Legacy StartBit fields present",
                        _sample(no_legacy)))

    orphan_bits = [k for k in db.get("J1939BitDecodings", {}) if k not in spns]
    checks.append(Check("Bit decodings", WARN if orphan_bits else PASS,
                        f"{len(orphan_bits)} bit decodings for undefined SPNs" if orphan_bits else "Bit decodings refer to defined SPNs",
                        _sample(orphan_bits)))

    declared = meta.get("units")
    wrong = []
    for num, spn in spns.items():
        s = table.system_of(spn.get("Units"))
        if declared and s and s != declared:
            wrong.append(f"SPN {num}: {spn.get('Units')}")
    system = unit_system(db, table)
    if declared:
        checks.append(Check("Unit system", FAIL if wrong else PASS,
                            f"Declared {declared}, but {len(wrong)} SPNs use other-system units" if wrong
                            else f"All convertible units are {declared}", _sample(wrong)))
    else:
        checks.append(Check("Unit system", INFO, f"Not declared; inferred {system or 'none'} from unit labels"))

    if not db.get("J1939SATabledb"):
        checks.append(Check("Source addresses", WARN, "Source address table is empty"))
    return checks


def compare(db, baseline, table=None):
    """Differences between `db` and `baseline` (another version or unit system)."""
    table = table or UnitTable()
    checks = []
    a_pgn, b_pgn = set(db["J1939PGNdb"]), set(baseline["J1939PGNdb"])
    a_spn, b_spn = set(db["J1939SPNdb"]), set(baseline["J1939SPNdb"])
    checks.append(Check("PGNs added", INFO, f"{len(a_pgn - b_pgn)} PGNs not in baseline", _sample(sorted(a_pgn - b_pgn, key=int))))
    checks.append(Check("PGNs removed", WARN if b_pgn - a_pgn else PASS, f"{len(b_pgn - a_pgn)} baseline PGNs missing",
                        _sample(sorted(b_pgn - a_pgn, key=int))))
    checks.append(Check("SPNs added", INFO, f"{len(a_spn - b_spn)} SPNs not in baseline", _sample(sorted(a_spn - b_spn, key=int))))
    checks.append(Check("SPNs removed", WARN if b_spn - a_spn else PASS, f"{len(b_spn - a_spn)} baseline SPNs missing",
                        _sample(sorted(b_spn - a_spn, key=int))))

    sys_a, sys_b = unit_system(db, table), unit_system(baseline, table)
    changed, conv_bad, conv_ok = [], [], 0
    for num in sorted(a_spn & b_spn, key=int):
        x, y = db["J1939SPNdb"][num], baseline["J1939SPNdb"][num]
        if x.get("SPNLength") != y.get("SPNLength"):
            changed.append(f"SPN {num}: length {y.get('SPNLength')} -> {x.get('SPNLength')}")
            continue
        rx, ry = _num(x.get("Resolution")), _num(y.get("Resolution"))
        ox, oy = _num(x.get("Offset")) or 0.0, _num(y.get("Offset")) or 0.0
        if rx is None or ry is None:
            continue
        if sys_a != sys_b and sys_a and sys_b:
            metric, us = (y, x) if sys_a == US else (x, y)
            conv = table.conversion_between(metric.get("Units"), us.get("Units"))
            if conv:
                rm, om = _num(metric.get("Resolution")), _num(metric.get("Offset")) or 0.0
                ru, ou = _num(us.get("Resolution")), _num(us.get("Offset")) or 0.0
                if not (_close(rm * conv["scale"], ru) and _close(om * conv["scale"] + conv["offset"], ou)):
                    conv_bad.append(f"SPN {num}: {metric.get('Units')} r={rm} o={om} vs {us.get('Units')} r={ru} o={ou}")
                else:
                    conv_ok += 1
                continue
        if not (_close(rx, ry) and _close(ox, oy)) or canonical_unit(x.get("Units")) != canonical_unit(y.get("Units")):
            changed.append(f"SPN {num}: {y.get('Resolution')}/{y.get('Offset')} {y.get('Units')} -> "
                           f"{x.get('Resolution')}/{x.get('Offset')} {x.get('Units')}")
    if sys_a != sys_b and sys_a and sys_b:
        checks.append(Check("Unit conversion", FAIL if conv_bad else PASS,
                            f"{conv_ok} converted SPNs consistent ({sys_b} baseline vs {sys_a})" if not conv_bad
                            else f"{len(conv_bad)} SPNs inconsistent with j1939_units.json", _sample(conv_bad)))
    checks.append(Check("Scaling changes", WARN if changed else PASS,
                        f"{len(changed)} common SPNs changed length, scaling or units" if changed else "Common SPNs unchanged",
                        _sample(changed)))
    return checks


def _close(a, b, rel=1e-6):
    if a is None or b is None:
        return False
    return abs(a - b) <= rel * max(1.0, abs(a), abs(b))


# --------------------------------------------------------------------------
# Cross-check against the Digital Annex SLOT definitions
# --------------------------------------------------------------------------

def _header_key(text):
    return re.sub(r"[^A-Z0-9()]+", "_", str(text or "").upper()).strip("_")


def _sheet_rows(path, sheet_name):
    """Yield rows (tuples) of one sheet from an .xlsx or .xls workbook."""
    if path.lower().endswith(".xls"):
        import xlrd
        book = xlrd.open_workbook(path, on_demand=True)
        sheet = book.sheet_by_name(sheet_name)
        for i in range(sheet.nrows):
            yield tuple(sheet.row_values(i))
        return
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for row in wb[sheet_name].iter_rows(values_only=True):
            yield row
    finally:
        wb.close()


def _sheet_names(path):
    if path.lower().endswith(".xls"):
        import xlrd
        return xlrd.open_workbook(path, on_demand=True).sheet_names()
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def _table(path, sheet_name, required):
    """Rows of a sheet as dicts keyed by normalized header, using the first row that has `required`."""
    rows = _sheet_rows(path, sheet_name)
    header = None
    for i, row in enumerate(rows):
        keys = [_header_key(c) for c in row]
        if required in keys:
            header = keys
            break
        if i > 15:
            break
    if header is None:
        return []
    out = []
    for row in rows:
        out.append({k: v for k, v in zip(header, row) if k})
    return out


def read_slots(da_paths):
    """SLOT definitions and the SPN -> SLOT map from Digital Annex workbook(s).

    Returns (slots, spn_slot, spn_row_values) where slots maps SLOT identifier to
    a dict(name, type, unit, transfer, scale, offset, range_max, len_min, len_max).
    """
    slots, spn_slot, spn_row = {}, {}, {}
    for path in da_paths:
        names = _sheet_names(path)
        if "SLOTs" in names:
            for r in _table(path, "SLOTs", "SLOT_IDENTIFIER"):
                sid = r.get("SLOT_IDENTIFIER")
                if sid in (None, ""):
                    continue
                slots[str(sid).strip()] = {
                    "name": r.get("SLOT_NAME"), "type": r.get("SLOT_TYPE"), "unit": r.get("UNIT") or "",
                    "transfer": str(r.get("TRANSFER_FUNCTION_TYPE") or "").strip().lower(),
                    "scale": _num(r.get("SCALE_FACTOR_(VALUE_ONLY)")),
                    "offset": _num(r.get("OFFSET_(VALUE_ONLY)")),
                    "range_max": _num(r.get("RANGE_MAXIMUM_(VALUE_ONLY)")),
                    "len_min": _num(r.get("LENGTH_MINIMUM_(BITS)")),
                    "len_max": _num(r.get("LENGTH_MAXIMUM_(BITS)")),
                }
        sheet = next((n for n in ("SPs & PGs", "SPNs & PGNs") if n in names), None)
        if sheet:
            for r in _table(path, sheet, "SLOT_IDENTIFIER"):
                spn, sid = r.get("SPN"), r.get("SLOT_IDENTIFIER")
                if spn in (None, "", "N/A") or sid in (None, ""):
                    continue
                try:
                    key = str(int(float(spn)))
                except (TypeError, ValueError):
                    continue
                spn_slot[key] = str(sid).strip()
                spn_row[key] = {"scale": _num(r.get("SCALE_FACTOR_(VALUE_ONLY)")),
                                "offset": _num(r.get("OFFSET_(VALUE_ONLY)")),
                                "unit": r.get("UNIT") or ""}
    return slots, spn_slot, spn_row


def _same_unit(a, b, table):
    ca, cb = canonical_unit(a), canonical_unit(b)
    if ca == cb:
        return True
    ka, kb = table.by_alias.get(ca), table.by_alias.get(cb)
    return ka is not None and ka == kb


def slot_crosscheck(da_paths, metric_db, us_db=None, table=None):
    """Check database scaling against the Digital Annex SLOT sheet.

    * DA consistency: each SP row's value-only scale/offset equals its SLOT's.
    * Metric database: Resolution, Offset and unit equal the SLOT values; SP length
      lies within the SLOT length limits.
    * US database: Resolution/Offset equal the SLOT values converted with
      j1939_units.json (or unchanged when the SLOT unit is not converted).
    * Summary of which SLOT units are converted to US units and which stay metric.
    """
    table = table or UnitTable()
    slots, spn_slot, spn_row = read_slots(da_paths)
    checks = []
    if not slots:
        return [Check("SLOT table", FAIL, "No SLOTs sheet found in the Digital Annex")]
    checks.append(Check("SLOT table", INFO, f"{len(slots)} SLOTs; {len(spn_slot)} SPNs reference a SLOT"))

    da_inconsistent = []
    for spn, sid in spn_slot.items():
        slot, row = slots.get(sid), spn_row[spn]
        if slot is None or slot["transfer"] != "numeric":
            continue
        if row["scale"] is not None and slot["scale"] is not None and not _close(row["scale"], slot["scale"], 1e-9):
            da_inconsistent.append(f"SPN {spn}: row scale {row['scale']} vs SLOT {sid} {slot['scale']}")
        elif row["offset"] is not None and slot["offset"] is not None and not _close(row["offset"], slot["offset"], 1e-9):
            da_inconsistent.append(f"SPN {spn}: row offset {row['offset']} vs SLOT {sid} {slot['offset']}")
    checks.append(Check("DA rows vs SLOT sheet", WARN if da_inconsistent else PASS,
                        f"{len(da_inconsistent)} SP rows disagree with their SLOT definition" if da_inconsistent
                        else "Every SP row matches its SLOT definition", _sample(da_inconsistent)))

    spns = metric_db.get("J1939SPNdb", {})
    missing_slot = [s for s in spns if s not in spn_slot]
    scale_bad, unit_bad, length_bad, checked = [], [], [], 0
    for spn, entry in spns.items():
        slot = slots.get(spn_slot.get(spn))
        if slot is None:
            continue
        length = entry.get("SPNLength")
        if isinstance(length, int) and slot["len_min"] is not None and slot["len_max"] is not None:
            if not slot["len_min"] <= length <= slot["len_max"]:
                length_bad.append(f"SPN {spn}: {length} bits outside SLOT {spn_slot[spn]} {slot['len_min']:g}..{slot['len_max']:g}")
        if slot["transfer"] != "numeric" or slot["scale"] is None:
            continue
        checked += 1
        res, off = _num(entry.get("Resolution")), _num(entry.get("Offset")) or 0.0
        if not (_close(res, slot["scale"], 1e-9) and _close(off, slot["offset"] or 0.0, 1e-9)):
            scale_bad.append(f"SPN {spn}: {res}/{off} vs SLOT {spn_slot[spn]} {slot['scale']}/{slot['offset']}")
        if slot["unit"] and not _same_unit(entry.get("Units"), slot["unit"], table):
            unit_bad.append(f"SPN {spn}: unit {entry.get('Units')!r} vs SLOT {spn_slot[spn]} {slot['unit']!r}")
    checks.append(Check("Metric scaling vs SLOT", FAIL if scale_bad else PASS,
                        f"{len(scale_bad)} of {checked} numeric SPNs differ from their SLOT" if scale_bad
                        else f"{checked} numeric SPNs match their SLOT scale and offset", _sample(scale_bad)))
    checks.append(Check("Metric units vs SLOT", WARN if unit_bad else PASS,
                        f"{len(unit_bad)} SPN units differ from their SLOT unit" if unit_bad
                        else "SPN units match their SLOT units", _sample(unit_bad)))
    checks.append(Check("Length vs SLOT limits", WARN if length_bad else PASS,
                        f"{len(length_bad)} SPN lengths outside their SLOT limits" if length_bad
                        else "SPN lengths are within SLOT limits", _sample(length_bad)))
    checks.append(Check("SPNs without SLOT", INFO, f"{len(missing_slot)} SPNs in the database have no SLOT reference",
                        _sample(sorted(missing_slot, key=int))))

    converted, kept = {}, {}
    for sid, slot in slots.items():
        if slot["transfer"] != "numeric" or not slot["unit"]:
            continue
        conv = table.lookup(slot["unit"])
        bucket = converted.setdefault(f'{conv["metric"]} -> {conv["us"]}', []) if conv else kept.setdefault(slot["unit"], [])
        bucket.append(sid)
    checks.append(Check("SLOT units converted to US", INFO, f"{sum(map(len, converted.values()))} numeric SLOTs in {len(converted)} units",
                        [f"{k}: {len(v)} SLOTs" for k, v in sorted(converted.items())]))
    checks.append(Check("SLOT units kept as published", INFO, f"{sum(map(len, kept.values()))} numeric SLOTs in {len(kept)} units "
                        "(add an entry to j1939_units.json to convert one)",
                        [f"{k}: {len(v)} SLOTs" for k, v in sorted(kept.items(), key=lambda kv: -len(kv[1]))]))

    if us_db is not None:
        us_spns = us_db.get("J1939SPNdb", {})
        us_bad, us_checked = [], 0
        for spn, entry in us_spns.items():
            slot = slots.get(spn_slot.get(spn))
            if slot is None or slot["transfer"] != "numeric" or slot["scale"] is None:
                continue
            us_checked += 1
            conv = table.lookup(slot["unit"]) if slot["unit"] else None
            scale, offset = slot["scale"], slot["offset"] or 0.0
            unit_ok = True
            if conv:
                scale, offset = scale * conv["scale"], offset * conv["scale"] + conv["offset"]
                unit_ok = canonical_unit(entry.get("Units")) == canonical_unit(conv["us"])
            res, off = _num(entry.get("Resolution")), _num(entry.get("Offset")) or 0.0
            if not (_close(res, scale, 1e-9) and _close(off, offset, 1e-9) and unit_ok):
                us_bad.append(f"SPN {spn}: {res}/{off} {entry.get('Units')} vs SLOT {spn_slot[spn]} "
                              f"converted {scale:.10g}/{offset:.10g} {conv['us'] if conv else slot['unit']}")
        checks.append(Check("US scaling vs SLOT conversion", FAIL if us_bad else PASS,
                            f"{len(us_bad)} of {us_checked} numeric SPNs differ from the converted SLOT" if us_bad
                            else f"{us_checked} numeric SPNs match their SLOT after US conversion", _sample(us_bad)))
    return checks


# --------------------------------------------------------------------------
# Test vectors
# --------------------------------------------------------------------------

VECTORS_DOC = [
    "Decode test vectors for J1939 databases. Each vector decodes one SPN from a PGN payload and checks",
    "'expected' (value per unit system, absolute 'tolerance'), 'expected_status' or 'expected_text'.",
    "LICENSING: vectors must not reproduce the SAE J1939 Digital Annex. They hold payloads and expected",
    "results only (no scaling, offsets, bit positions, ranges or DA descriptions), cover at most 40",
    "widely published SPNs with at most 3 vectors each, and use our own names. tests/test_no_licensed_content.py",
    "enforces this. Regenerate with: python tests/build_vectors.py (requires the local licensed databases).",
]


def load_vectors(path=VECTORS_FILE):
    with open(path, encoding="utf-8") as f:
        return json.load(f)["vectors"]


def save_vectors(vectors, path=VECTORS_FILE):
    with open(path, "w", encoding="utf-8") as f:
        json.dump({"_doc": VECTORS_DOC, "vectors": vectors}, f, indent=1)


def run_vector(db, vector, table=None, system=None):
    """Returns (status, message). status is PASS, FAIL or SKIP.

    A vector checks one of: `expected_status` (e.g. NotAvailable, Error),
    `expected_text` (ASCII fields), or `expected` values per unit system.
    """
    table = table or UnitTable()
    data = bytes.fromhex(str(vector["data"]).replace(" ", ""))
    result = decode_spn(db, int(vector["pgn"]), int(vector["spn"]), data)
    if result is None:
        return "SKIP", f"PGN {vector['pgn']} / SPN {vector['spn']} not in database"
    if vector.get("expected_status"):
        ok = result["status"] == vector["expected_status"]
        return ("PASS" if ok else "FAIL"), f"status {result['status']} (expected {vector['expected_status']})"
    if vector.get("expected_text") is not None:
        ok = result["text"] == vector["expected_text"]
        return ("PASS" if ok else "FAIL"), f"text {result['text']!r} (expected {vector['expected_text']!r})"
    if result["value"] is None:
        return "FAIL", f"not decodable ({result['status']})"
    if result["status"] != "Valid":
        return "FAIL", f"status {result['status']}, expected a valid value"
    unit_sys = table.system_of(result["units"]) or system or unit_system(db, table) or METRIC
    expected = vector.get("expected", {})
    want = _num(expected.get(unit_sys, expected.get(METRIC)))
    if want is None:
        return "SKIP", f"no expected value for {unit_sys}"
    tol = _num(vector.get("tolerance")) or 1e-6
    ok = abs(result["value"] - want) <= tol
    return ("PASS" if ok else "FAIL"), f"{result['value']:g} {result['units']} (expected {want:g}, {unit_sys})"


def run_vectors(db, vectors, table=None):
    table = table or UnitTable()
    return [(v, *run_vector(db, v, table)) for v in vectors]


# --------------------------------------------------------------------------
# Command line
# --------------------------------------------------------------------------

def _print_checks(title, checks):
    print(f"\n== {title}")
    for c in checks:
        print(f"[{c.level}] {c.name}: {c.message}")
        for d in c.details:
            print(f"        {d}")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    g = sub.add_parser("generate", help="build licensed databases from Digital Annex files")
    g.add_argument("da", nargs="+")
    g.add_argument("--out", default=".")
    g.add_argument("--units", choices=["both", METRIC, US], default="both")
    v = sub.add_parser("validate", help="check a database file")
    v.add_argument("db")
    v.add_argument("--baseline")
    v.add_argument("--vectors", default=VECTORS_FILE)
    v.add_argument("--da", nargs="+", help="Digital Annex workbook(s) for the SLOT cross-check")
    v.add_argument("--us", help="US customary database to check against the SLOTs (with --da)")
    args = ap.parse_args(argv)

    if args.cmd == "generate":
        systems = (METRIC, US) if args.units == "both" else (args.units,)
        generate(args.da, args.out, systems)
        return 0
    db = load(args.db)
    table = UnitTable()
    checks = validate(db, table)
    _print_checks(f"Validation of {args.db}", checks)
    if args.baseline:
        cmp = compare(db, load(args.baseline), table)
        _print_checks(f"Comparison with {args.baseline}", cmp)
        checks += cmp
    if args.da:
        metric_db, us_db = db, (load(args.us) if args.us else None)
        if unit_system(db, table) == US:
            metric_db, us_db = (load(args.baseline) if args.baseline else None), db
        if metric_db is not None:
            slot_checks = slot_crosscheck(args.da, metric_db, us_db, table)
            _print_checks("Cross-check with Digital Annex SLOTs", slot_checks)
            checks += slot_checks
    failed = sum(c.level == FAIL for c in checks)
    if args.vectors and os.path.exists(args.vectors):
        print(f"\n== Test vectors ({args.vectors})")
        for vec, status, msg in run_vectors(db, load_vectors(args.vectors), table):
            print(f"[{status}] {vec.get('name')}: {msg}")
            failed += status == "FAIL"
    print(f"\n{failed} failure(s)")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
